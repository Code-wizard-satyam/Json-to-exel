"""Shared JSON inspection and Excel export logic for the Streamlit app and API."""

from __future__ import annotations

import io
import json
from collections import OrderedDict
from collections.abc import Mapping, Sequence
from datetime import datetime, timezone
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_UPLOAD_BYTES = 25 * 1024 * 1024


def parse_json(raw: bytes | str) -> Any:
    """Parse JSON bytes and return a native Python value."""
    if isinstance(raw, bytes):
        raw = raw.decode("utf-8-sig")
    return json.loads(raw)


def _display_value(value: Any) -> Any:
    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        return json.dumps(
            value,
            ensure_ascii=False,
            separators=(",", ":"),
        )

    return value


def _walk_paths(
    value: Any,
    path: str,
    rows: list[dict[str, Any]],
) -> None:
    if isinstance(value, Mapping):
        if not value:
            rows.append(
                {
                    "path": path or "$",
                    "value": "{}",
                }
            )
            return

        for key, child in value.items():
            child_path = f"{path}.{key}" if path else str(key)
            _walk_paths(child, child_path, rows)

        return

    if isinstance(value, Sequence) and not isinstance(
        value,
        (str, bytes, bytearray),
    ):
        if not value:
            rows.append(
                {
                    "path": path or "$",
                    "value": "[]",
                }
            )
            return

        for index, child in enumerate(value):
            _walk_paths(
                child,
                f"{path}[{index}]",
                rows,
            )

        return

    rows.append(
        {
            "path": path or "$",
            "value": _display_value(value),
        }
    )


def path_rows(data: Any) -> list[dict[str, Any]]:
    """Return one row per leaf value, preserving the full JSON path."""
    rows: list[dict[str, Any]] = []
    _walk_paths(data, "", rows)
    return rows


def _find_camera_maps(
    value: Any,
    path: str = "",
) -> list[tuple[str, Mapping[str, Any]]]:
    found: list[tuple[str, Mapping[str, Any]]] = []

    if isinstance(value, Mapping):
        for key, child in value.items():
            child_path = f"{path}.{key}" if path else str(key)

            if key == "cameras" and isinstance(child, Mapping):
                found.append((child_path, child))

            found.extend(
                _find_camera_maps(
                    child,
                    child_path,
                )
            )

    elif isinstance(value, Sequence) and not isinstance(
        value,
        (str, bytes, bytearray),
    ):
        for index, child in enumerate(value):
            found.extend(
                _find_camera_maps(
                    child,
                    f"{path}[{index}]",
                )
            )

    return found


def alarm_rows(data: Any) -> list[dict[str, Any]]:
    """Extract camera alarm objects into a useful tabular shape."""
    rows: list[dict[str, Any]] = []

    for camera_path, cameras in _find_camera_maps(data):
        endpoint = camera_path.removesuffix(".cameras")

        for camera_id, camera in cameras.items():
            if not isinstance(camera, Mapping):
                rows.append(
                    {
                        "endpoint": endpoint,
                        "camera_id": camera_id,
                        "category": "value",
                        "level": "",
                        "text": _display_value(camera),
                    }
                )
                continue

            for category, details in camera.items():
                detail_items = (
                    details
                    if isinstance(details, Sequence)
                    and not isinstance(
                        details,
                        (str, bytes, bytearray),
                    )
                    else [details]
                )

                for detail in detail_items:
                    if isinstance(detail, Mapping):
                        rows.append(
                            {
                                "endpoint": endpoint,
                                "camera_id": camera_id,
                                "category": category,
                                "level": detail.get("level", ""),
                                "text": detail.get("text", ""),
                            }
                        )
                    else:
                        rows.append(
                            {
                                "endpoint": endpoint,
                                "camera_id": camera_id,
                                "category": category,
                                "level": "",
                                "text": _display_value(detail),
                            }
                        )

    return rows


def _record_rows(
    value: Any,
    context: dict[str, Any],
    rows: list[dict[str, Any]],
) -> None:
    """Flatten a nested value into records with useful identifiers."""
    if isinstance(value, Mapping):
        if not value:
            rows.append(
                {
                    **context,
                    "value": "{}",
                }
            )
            return

        scalar_items = {
            str(key): _display_value(child)
            for key, child in value.items()
            if not isinstance(
                child,
                (Mapping, list, tuple),
            )
        }

        nested_items = {
            key: child
            for key, child in value.items()
            if isinstance(
                child,
                (Mapping, list, tuple),
            )
        }

        if scalar_items:
            rows.append(
                {
                    **context,
                    **scalar_items,
                }
            )

        for key, child in nested_items.items():
            child_context = {
                **context,
                "group": str(key),
            }

            _record_rows(
                child,
                child_context,
                rows,
            )

        return

    if isinstance(value, Sequence) and not isinstance(
        value,
        (str, bytes, bytearray),
    ):
        if not value:
            rows.append(
                {
                    **context,
                    "value": "[]",
                }
            )
            return

        for index, child in enumerate(value):
            _record_rows(
                child,
                {
                    **context,
                    "index": index,
                },
                rows,
            )

        return

    rows.append(
        {
            **context,
            "value": _display_value(value),
        }
    )


def record_rows(data: Any) -> list[dict[str, Any]]:
    """Create generalized rows from metadata and report endpoints."""
    rows: list[dict[str, Any]] = []

    if not isinstance(data, Mapping):
        _record_rows(
            data,
            {"section": "root"},
            rows,
        )
        return rows

    for key, value in data.items():
        if key == "reply" and isinstance(value, Mapping):
            for endpoint, endpoint_value in value.items():
                payload = (
                    endpoint_value.get("reply")
                    if isinstance(endpoint_value, Mapping)
                    and "reply" in endpoint_value
                    else endpoint_value
                )

                _record_rows(
                    payload,
                    {"endpoint": str(endpoint)},
                    rows,
                )

        elif not isinstance(
            value,
            (Mapping, list, tuple),
        ):
            rows.append(
                {
                    "section": "metadata",
                    "field": str(key),
                    "value": _display_value(value),
                }
            )

    if not rows:
        _record_rows(
            data,
            {"section": "root"},
            rows,
        )

    return rows


def build_tables(
    data: Any,
    source_name: str,
) -> OrderedDict[str, pd.DataFrame]:
    """Build named DataFrames that become workbook sheets."""
    tables: OrderedDict[str, pd.DataFrame] = OrderedDict()

    rows = record_rows(data)
    alarms = alarm_rows(data)
    paths = path_rows(data)

    endpoint_count = 0

    if isinstance(data, Mapping) and isinstance(
        data.get("reply"),
        Mapping,
    ):
        endpoint_count = len(data["reply"])

    summary = [
        {
            "metric": "Source file",
            "value": source_name,
        },
        {
            "metric": "Generated at (UTC)",
            "value": datetime.now(timezone.utc).isoformat(),
        },
        {
            "metric": "JSON root type",
            "value": type(data).__name__,
        },
        {
            "metric": "Top-level keys",
            "value": len(data) if isinstance(data, Mapping) else "",
        },
        {
            "metric": "Report endpoints",
            "value": endpoint_count,
        },
        {
            "metric": "Generalized records",
            "value": len(rows),
        },
        {
            "metric": "Camera alarm rows",
            "value": len(alarms),
        },
        {
            "metric": "Leaf JSON paths",
            "value": len(paths),
        },
    ]

    tables["Summary"] = pd.DataFrame(summary)

    if alarms:
        tables["Alarms"] = pd.DataFrame(alarms)

    if rows:
        tables["Records"] = pd.DataFrame(rows)

    if paths:
        tables["JSON Paths"] = pd.DataFrame(paths)

    return tables


def _safe_sheet_name(
    name: str,
    used: set[str],
) -> str:
    invalid_chars = '[]:*?/"\\'

    candidate = "".join(
        "_" if char in invalid_chars else char
        for char in name
    )[:31] or "Sheet"

    base = candidate
    counter = 2

    while candidate in used:
        suffix = f" {counter}"
        candidate = (
            f"{base[:31 - len(suffix)]}{suffix}"
        )
        counter += 1

    used.add(candidate)
    return candidate


def workbook_bytes(
    tables: Mapping[str, pd.DataFrame],
) -> bytes:
    """Render tables to a styled XLSX workbook."""
    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        used_names: set[str] = set()

        for sheet_name, frame in tables.items():
            safe_name = _safe_sheet_name(
                sheet_name,
                used_names,
            )

            frame.to_excel(
                writer,
                index=False,
                sheet_name=safe_name,
            )

            worksheet = writer.book[safe_name]

            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.row_dimensions[1].height = 28

            for cell in worksheet[1]:
                cell.font = Font(
                    bold=True,
                    color="FFFFFF",
                )
                cell.fill = PatternFill(
                    "solid",
                    fgColor="183B56",
                )
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                )

            for column_index, column_cells in enumerate(
                worksheet.iter_cols(
                    min_row=1,
                    max_row=min(
                        worksheet.max_row,
                        200,
                    ),
                ),
                start=1,
            ):
                longest = max(
                    (
                        len(str(cell.value))
                        if cell.value is not None
                        else 0
                        for cell in column_cells
                    ),
                    default=0,
                )

                worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = min(
                    max(longest + 2, 12),
                    52,
                )

    output.seek(0)
    return output.getvalue()